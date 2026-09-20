from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"C:\Users\Thien\Downloads\corpus")
SOURCE = ROOT / "PITCH DECK.pptx"
OUTPUT = ROOT / "PITCH DECK - MEMORY WITH A RECEIPT - FINAL.pptx"

TEAL = RGBColor(9, 149, 169)
DEEP_TEAL = RGBColor(0, 92, 128)
BLUE = RGBColor(22, 77, 113)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(18, 26, 34)
MUTED = RGBColor(103, 113, 121)
PALE = RGBColor(241, 246, 248)
LIGHT_TEAL = RGBColor(218, 241, 245)


def clear_text(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            clear_text(child)
    elif getattr(shape, "has_text_frame", False):
        shape.text_frame.clear()


def clear_slide_text(slide):
    for shape in slide.shapes:
        clear_text(shape)


def remove_text_placeholders(slide):
    """Remove empty template prompts while retaining any template photo slot."""
    for shape in list(slide.shapes):
        if not getattr(shape, "is_placeholder", False):
            continue
        if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
            continue
        shape.element.getparent().remove(shape.element)


def delete_slide(prs, index):
    slide_id = prs.slides._sldIdLst[index]
    r_id = slide_id.rId
    prs.part.drop_rel(r_id)
    del prs.slides._sldIdLst[index]


def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False,
            font="Arial", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.02, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return shape


def paragraphs(slide, x, y, w, h, rows, margin=0.02,
               align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, row in enumerate(rows):
        text, size, color, bold, after = row
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        run = p.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, transparency=0, line=None,
         radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line if line is not None else fill
    return shape


def line(slide, x1, y1, x2, y2, color=TEAL, width=2.25):
    shape = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def circle(slide, x, y, d, fill=TEAL, line_color=TEAL, text=None,
           text_color=WHITE, text_size=16):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line_color
    if text is not None:
        tf = shape.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(text_size)
        run.font.bold = True
        run.font.color.rgb = text_color
    return shape


def footer(slide, number, dark=False):
    color = WHITE if dark else MUTED
    textbox(slide, 0.95, 7.05, 2.6, 0.2, "AALTOAI HACKATHON 2026", 8, color)
    textbox(slide, 5.15, 7.05, 3.1, 0.2, "MEMORY WITH A RECEIPT", 8, color, align=PP_ALIGN.CENTER)
    textbox(slide, 12.55, 7.05, 0.25, 0.2, str(number), 8, color, align=PP_ALIGN.RIGHT)


def title(slide, text, x=1.0, y=0.62, w=11.5, color=BLACK, size=32,
          align=PP_ALIGN.LEFT):
    textbox(slide, x, y, w, 0.55, text, size, color, True, align=align)


def card(slide, x, y, w, h, heading, body, dark=False, heading_size=16,
         body_size=13):
    fill = RGBColor(255, 255, 255) if dark else WHITE
    text_color = INK if dark else INK
    rect(slide, x, y, w, h, fill, transparency=8 if dark else 0,
         line=fill, radius=True)
    rect(slide, x, y, w, 0.10, TEAL, line=TEAL)
    textbox(slide, x + 0.18, y + 0.25, w - 0.36, 0.35, heading,
            heading_size, BLUE, True, align=PP_ALIGN.CENTER)
    textbox(slide, x + 0.22, y + 0.82, w - 0.44, h - 1.0, body,
            body_size, text_color, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
            margin=0.01)


def add_pair(slide, y, left, right, x_left=3.55, x_right=7.45):
    textbox(slide, x_left, y, 3.2, 0.30, left, 12, MUTED, True)
    textbox(slide, x_right, y, 4.7, 0.30, right, 12, BLUE, True)
    line(slide, x_left, y + 0.38, 12.15, y + 0.38, RGBColor(205, 214, 218), 1.0)


def build():
    prs = Presentation(str(SOURCE))
    keep = [0, 2, 3, 11, 8, 12, 7, 10, 14, 20]
    original_ids = list(prs.slides._sldIdLst)
    for idx in range(len(prs.slides) - 1, -1, -1):
        if idx not in keep:
            delete_slide(prs, idx)
    desired_ids = [original_ids[idx] for idx in keep]
    for child in desired_ids:
        prs.slides._sldIdLst.remove(child)
    for child in desired_ids:
        prs.slides._sldIdLst.append(child)
    for slide in prs.slides:
        clear_slide_text(slide)
        remove_text_placeholders(slide)

    # 1. Cover
    slide = prs.slides[0]
    textbox(slide, 1.0, 4.82, 8.7, 0.62, "MEMORY WITH A RECEIPT", 34, WHITE, True)
    textbox(slide, 1.02, 5.46, 8.2, 0.42, "Organizational Memory Auditor", 20, WHITE)
    textbox(slide, 1.03, 6.18, 7.0, 0.25,
            "Every answer carries the evidence, the conflict, and the way back.",
            11, WHITE, italic=True)
    textbox(slide, 1.03, 6.53, 5.4, 0.22, "RELEX Challenge · AaltoAI Hackathon 2026", 9, WHITE)

    # 2. The challenge
    slide = prs.slides[1]
    title(slide, "WHY MEMORY BREAKS", 1.0, 0.60, 10.0, WHITE, 32)
    textbox(slide, 1.0, 1.24, 8.7, 0.38,
            "A clean summary can still be wrong when the archive disagrees with itself.",
            16, WHITE)
    paragraphs(slide, 1.10, 1.85, 5.0, 3.95, [
        ("45 files, one messy truth", 17, BLUE, True, 3),
        ("Transcripts, email threads and status reports — with people, dates and decisions changing over time.", 13, INK, False, 14),
        ("Suggestion ≠ commitment", 17, BLUE, True, 3),
        ("A polished assistant can turn an idea into an agreement if nobody checks the receipt.", 13, INK, False, 14),
    ])
    paragraphs(slide, 6.28, 1.85, 5.0, 3.95, [
        ("Stale ≠ false", 17, BLUE, True, 3),
        ("A later report may be superseded by operational evidence — or contradict it. Newer is not automatically truer.", 13, INK, False, 14),
        ("Privacy ≠ amnesia", 17, BLUE, True, 3),
        ("Remove a person from ordinary surfaces without deleting the organization's attributable history.", 13, INK, False, 14),
    ])
    rect(slide, 0.85, 6.12, 11.35, 0.58, DEEP_TEAL, line=DEEP_TEAL)
    textbox(slide, 1.08, 6.25, 10.8, 0.28,
            "The challenge: answer live questions with a receipt — then survive the privacy test.",
            14, WHITE, True)
    footer(slide, 2, dark=True)

    # 3. The product answer
    slide = prs.slides[2]
    title(slide, "THE ANSWER IS A CASE — NOT A CHAT", 6.78, 0.64, 6.05, BLACK, 22)
    textbox(slide, 6.82, 1.30, 5.55, 0.38,
            "Every question becomes a structured ruling.",
            15, BLUE, True)
    circle(slide, 6.84, 2.08, 0.36, TEAL, TEAL, "1", WHITE, 12)
    paragraphs(slide, 7.34, 2.03, 4.85, 0.68, [
        ("ASK", 13, BLUE, True, 0),
        ("Hybrid search finds the right moments — not the whole archive.", 12, INK, False, 0),
    ])
    circle(slide, 6.84, 3.00, 0.36, TEAL, TEAL, "2", WHITE, 12)
    paragraphs(slide, 7.34, 2.95, 4.85, 0.68, [
        ("CHECK", 13, BLUE, True, 0),
        ("Risk rules decide when the first answer needs a real challenge.", 12, INK, False, 0),
    ])
    circle(slide, 6.84, 3.92, 0.36, TEAL, TEAL, "3", WHITE, 12)
    paragraphs(slide, 7.34, 3.87, 4.85, 0.68, [
        ("RECONCILE", 13, BLUE, True, 0),
        ("Conflicts, uncertainty and Decision Evolution stay visible.", 12, INK, False, 0),
    ])
    circle(slide, 6.84, 4.84, 0.36, TEAL, TEAL, "4", WHITE, 12)
    paragraphs(slide, 7.34, 4.79, 4.85, 0.68, [
        ("RECEIPT", 13, BLUE, True, 0),
        ("Claims link to exact stored evidence — metadata comes from the locker, not the model.", 12, INK, False, 0),
    ])
    textbox(slide, 6.84, 5.92, 5.2, 0.46,
            "The LLM interprets. The Evidence Locker remembers.",
            15, TEAL, True, italic=True)
    footer(slide, 3)

    # 4. Roadmap and delivery
    slide = prs.slides[3]
    title(slide, "ROADMAP & SEAMLESS EXECUTION", 1.0, 0.45, 11.2, WHITE, 31,
          align=PP_ALIGN.CENTER)
    textbox(slide, 2.0, 1.12, 9.3, 0.35,
            "We planned for evidence first — and shipped the whole loop in that order.",
            16, WHITE, align=PP_ALIGN.CENTER)
    card(slide, 1.55, 2.05, 3.25, 2.75, "1 · LOCK THE MEMORY",
         "45 source files\n→ 2,534 atomic Evidence Units\n→ stable, citable IDs", dark=True)
    card(slide, 5.02, 2.05, 3.25, 2.75, "2 · BUILD THE AUDITOR",
         "Retrieve\n→ challenge\n→ reconcile\n→ validate", dark=True)
    card(slide, 8.49, 2.05, 3.25, 2.75, "3 · PROVE THE EDGES",
         "490 passing tests\n+ privacy/reversal rehearsal\n+ live demo path", dark=True)
    textbox(slide, 2.15, 5.34, 9.0, 0.38,
            "PLAN  →  BUILD  →  VERIFY  →  DEMO", 17, WHITE, True,
            align=PP_ALIGN.CENTER)
    textbox(slide, 2.0, 5.86, 9.3, 0.34,
            "What the judge sees: one coherent product, not a collection of promises.",
            13, WHITE, italic=True, align=PP_ALIGN.CENTER)
    footer(slide, 4, dark=True)

    # 5. Trust boundary and evidence locker
    slide = prs.slides[4]
    title(slide, "THE EVIDENCE LOCKER IS THE MEMORY", 1.0, 0.58, 11.2, WHITE, 30,
          align=PP_ALIGN.CENTER)
    textbox(slide, 2.0, 1.16, 9.3, 0.35,
            "Raw evidence is authoritative; model output is a derived interpretation.",
            16, WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, 1.10, 2.02, 3.6, 0.72, "45", 37, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.08, 2.74, 3.65, 0.45, "source documents", 15, WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, 4.86, 2.02, 3.6, 0.72, "2,534", 37, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 4.84, 2.74, 3.65, 0.45, "citable evidence units", 15, WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, 8.62, 2.02, 3.6, 0.72, "490", 37, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 8.60, 2.74, 3.65, 0.45, "automated tests", 15, WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, 1.28, 3.63, 3.3, 0.42, "PROPOSE", 13, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 4.62, 3.63, 4.1, 0.42, "CHALLENGE + RECONCILE", 13, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 9.05, 3.63, 3.0, 0.42, "VALIDATE", 13, WHITE, True, align=PP_ALIGN.CENTER)
    line(slide, 4.62, 3.84, 4.38, 3.84, TEAL, 3)
    line(slide, 8.98, 3.84, 8.75, 3.84, TEAL, 3)
    textbox(slide, 1.30, 4.55, 10.75, 0.36,
            "Decision Evolution keeps the path visible: proposed → agreed → contradicted → superseded → current.",
            15, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.68, 5.05, 10.0, 0.36,
            "No citation, no claim. No valid citation, no screen time.",
            17, WHITE, True, italic=True, align=PP_ALIGN.CENTER)
    footer(slide, 5, dark=True)

    # 6. Radar initiative
    slide = prs.slides[5]
    rect(slide, 0.75, 1.45, 11.85, 5.2, WHITE, line=WHITE)
    title(slide, "THE RADAR: INITIATIVE THAT LOOKS AHEAD", 1.0, 0.60, 11.6, BLACK, 28)
    textbox(slide, 1.0, 1.08, 11.2, 0.30,
            "A proactive feature that turns hindsight into a ranked reason to look again.",
            15, MUTED)
    nodes = [
        (1.15, "REJECTED /\nDEFERRED IDEA", "What did we stop?"),
        (4.05, "ORIGINAL\nBLOCKER", "Why did it stop?"),
        (6.95, "WHAT\nCHANGED?", "Does the blocker still hold?"),
        (9.85, "RADAR\nASSESSMENT", "Revisit · blocked · unclear"),
    ]
    for x, head, body in nodes:
        circle(slide, x + 0.72, 2.28, 0.52, TEAL, TEAL)
        textbox(slide, x, 2.94, 2.0, 0.55, head, 13, BLUE, True, align=PP_ALIGN.CENTER)
        textbox(slide, x - 0.06, 3.62, 2.12, 0.55, body, 11, INK, align=PP_ALIGN.CENTER)
    for x1, x2 in [(2.05, 4.02), (4.95, 6.92), (7.85, 9.82)]:
        line(slide, x1, 2.54, x2, 2.54, TEAL, 2.5)
    rect(slide, 1.25, 4.73, 10.9, 0.83, PALE, line=PALE, radius=True)
    textbox(slide, 1.60, 4.93, 10.2, 0.30,
            "7 skeptical checks keep hindsight honest — and the Radar never pretends it is evidence.",
            14, BLUE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 2.0, 5.90, 9.3, 0.36,
            "No prompt required. No unsupported recommendation. Just the next question worth asking.",
            15, TEAL, True, italic=True, align=PP_ALIGN.CENTER)
    footer(slide, 6)

    # 7. Reversible pseudonymisation
    slide = prs.slides[6]
    title(slide, "REVERSIBLE PSEUDONYMIZATION", 1.0, 0.55, 11.2, WHITE, 30,
          align=PP_ALIGN.CENTER)
    textbox(slide, 2.0, 1.10, 9.3, 0.35,
            "Privacy without erasing organizational memory.", 17, WHITE,
            align=PP_ALIGN.CENTER)
    card(slide, 1.55, 2.04, 3.25, 2.88, "PRESERVE",
         "Every evidence unit, relationship, receipt and profile survives.", dark=True)
    card(slide, 5.02, 2.04, 3.25, 2.88, "PROTECT",
         "Names become a stable alias — Participant Q7M4-N8 — while the original identity stays in an encrypted reversal vault.", dark=True,
         body_size=12)
    card(slide, 8.49, 2.04, 3.25, 2.88, "RESTORE",
         "An authenticated admin can reverse it, rebuild the indexes, and verify every surface.", dark=True)
    textbox(slide, 1.72, 5.45, 9.95, 0.42,
            "Source · search · embeddings · cases · cache: rebuilt, rechecked, and never filtered at query time.",
            14, BLUE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 2.0, 6.10, 9.3, 0.32,
            "The person changes name in the product — not history in the archive.",
            15, WHITE, italic=True, align=PP_ALIGN.CENTER)
    footer(slide, 7, dark=True)

    # 8. Differentiators
    slide = prs.slides[7]
    title(slide, "NOT JUST ANOTHER CHATBOT", 1.0, 0.58, 11.2, WHITE, 31,
          align=PP_ALIGN.CENTER)
    textbox(slide, 1.55, 1.18, 10.3, 0.32,
            "The difference is accountable memory — not more fluent prose.",
            16, WHITE, align=PP_ALIGN.CENTER)
    textbox(slide, 3.55, 1.95, 3.2, 0.34, "TYPICAL CHATBOT", 14, MUTED, True)
    textbox(slide, 7.45, 1.95, 4.7, 0.34, "OUR AUDITOR", 14, BLUE, True)
    add_pair(slide, 2.48, "Best-effort prose", "Typed Case with a status")
    add_pair(slide, 3.24, "One answer", "Evidence + objections + uncertainty")
    add_pair(slide, 4.00, "Latest statement wins", "Decision Evolution shows supersession")
    add_pair(slide, 4.76, "Forget by filtering", "Pseudonymise, preserve, verify")
    add_pair(slide, 5.52, "Waits for a question", "Radar surfaces what to revisit")
    textbox(slide, 3.35, 6.28, 8.6, 0.34,
            "A chatbot answers. An auditor can defend the answer.",
            16, WHITE, True, italic=True, align=PP_ALIGN.CENTER)
    footer(slide, 8, dark=True)

    # 9. Live demo story
    slide = prs.slides[8]
    rect(slide, 0.55, 1.18, 12.1, 5.4, WHITE, line=WHITE)
    title(slide, "THE LIVE STORY IN 60 SECONDS", 1.0, 0.58, 11.2, BLACK, 31)
    textbox(slide, 1.0, 1.08, 11.2, 0.31,
            "One question. One evidence trail. One controlled way to forget responsibly.",
            15, MUTED)
    steps = [
        (1.0, "ASK", "Did Acme sign off UAT?"),
        (3.05, "OPEN", "Exact turn in context"),
        (5.10, "EVOLVE", "Proposal → current state"),
        (7.15, "RADAR", "What should we revisit?"),
        (9.20, "ALIAS", "Pseudonymise one person"),
        (11.25, "RE-ASK", "History intact"),
    ]
    line(slide, 1.22, 3.08, 12.15, 3.08, TEAL, 3)
    for idx, (x, head, body) in enumerate(steps, 1):
        circle(slide, x, 2.77, 0.62, TEAL, TEAL, str(idx), WHITE, 15)
        textbox(slide, x - 0.40, 3.55, 1.43, 0.30, head, 12, BLUE, True, align=PP_ALIGN.CENTER)
        textbox(slide, x - 0.66, 4.02, 1.95, 0.68, body, 11, INK, align=PP_ALIGN.CENTER)
    rect(slide, 2.0, 5.45, 9.4, 0.72, PALE, line=PALE, radius=True)
    textbox(slide, 2.28, 5.68, 8.85, 0.25,
            "Same archive. New answer. Zero amnesia.", 16, TEAL, True,
            align=PP_ALIGN.CENTER)
    footer(slide, 9)

    # 10. Close
    slide = prs.slides[9]
    textbox(slide, 0.98, 4.06, 5.65, 0.62,
            "WE BUILT MEMORY THAT CAN EXPLAIN ITSELF — AND FORGET RESPONSIBLY.",
            24, WHITE, True)
    textbox(slide, 7.22, 4.10, 4.95, 1.12,
            "A receipt for every conclusion.\nA controlled way back.\nA product that does more than chat.",
            18, WHITE)
    textbox(slide, 7.22, 5.52, 4.95, 0.35,
            "Thank you · Questions welcome", 15, TEAL, True)
    textbox(slide, 1.02, 6.15, 5.2, 0.28,
            "RELEX Challenge · AaltoAI Hackathon 2026", 10, MUTED)

    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
