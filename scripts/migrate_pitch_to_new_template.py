from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"C:\Users\Thien\Downloads\corpus")
TEMPLATE = Path(r"C:\Users\Thien\Documents\Calc1\pitch .pptx")
OUTPUT = ROOT / "PITCH DECK - MEMORY WITH A RECEIPT - NEW TEMPLATE.pptx"

FONT = "Arial"
BLACK = RGBColor(0, 0, 0)
TEAL = RGBColor(9, 149, 169)
BLUE = RGBColor(22, 77, 113)
MUTED = RGBColor(96, 103, 110)


def remove_template_text(slide):
    """Keep template imagery and tables, remove editable text prompts/content."""
    for shape in list(slide.shapes):
        if not getattr(shape, "is_placeholder", False):
            continue
        if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            continue
        shape.element.getparent().remove(shape.element)


def text_box(slide, x, y, w, h, text, size=14, color=BLACK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, italic=False,
             margin=0.02):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return shape


def rich_box(slide, x, y, w, h, rows, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, margin=0.02):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, (text, size, color, bold, after) in enumerate(rows):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def metadata(slide, number):
    text_box(slide, 0.15, 0.12, 3.82, 0.25,
             "RELEX CHALLENGE · 2026", 9, MUTED)
    text_box(slide, 6.53, 0.12, 5.8, 0.25,
             "MEMORY WITH A RECEIPT", 9, MUTED, align=PP_ALIGN.CENTER)
    text_box(slide, 12.72, 0.12, 0.47, 0.25,
             str(number), 9, MUTED, align=PP_ALIGN.RIGHT)


def set_cell(cell, text, size=14, bold=False, color=BLACK):
    tf = cell.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def find_table(slide):
    return next(shape.table for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.TABLE)


def build():
    prs = Presentation(str(TEMPLATE))
    if len(prs.slides) != 10:
        raise ValueError(f"Expected 10 template slides, found {len(prs.slides)}")
    for slide in prs.slides:
        remove_template_text(slide)

    # 1. Cover — keep the template's landscape image and editorial title treatment.
    slide = prs.slides[0]
    text_box(slide, 0.25, 4.62, 5.3, 0.30,
             "ENGRAM", 11, MUTED, True)
    text_box(slide, 5.90, 4.62, 3.8, 0.30,
             "AALTOAI HACKATHON 2026", 11, MUTED, align=PP_ALIGN.CENTER)
    text_box(slide, 0.20, 5.50, 10.4, 1.85,
             "Memory with a\nReceipt", 48, BLACK, True)

    # 2. The challenge — use the welcome split layout.
    slide = prs.slides[1]
    metadata(slide, 2)
    text_box(slide, 0.18, 1.72, 5.75, 3.0,
             "Why memory\nbreaks", 40, BLACK, True)
    rich_box(slide, 6.53, 1.66, 6.10, 3.10, [
        ("A clean summary can still be wrong when the archive disagrees with itself.", 16, BLACK, True, 16),
        ("45 files, one messy truth", 14, BLUE, True, 2),
        ("Suggestion ≠ commitment", 14, BLUE, True, 2),
        ("Stale ≠ false", 14, BLUE, True, 2),
        ("Privacy ≠ amnesia", 14, BLUE, True, 0),
    ])

    # 3. Archive scale — preserve the large-stat template.
    slide = prs.slides[2]
    metadata(slide, 3)
    text_box(slide, 0.18, 1.35, 5.75, 2.0,
             "2,534", 78, BLACK, True)
    text_box(slide, 0.25, 3.60, 5.3, 0.40,
             "citable Evidence Units", 19, BLUE, True)
    text_box(slide, 6.67, 1.45, 5.55, 2.7,
             "from 45 source\ndocuments", 36, BLACK, True)
    text_box(slide, 6.70, 4.55, 5.1, 0.80,
             "March 2024 – July 2026\n490 automated tests", 18, MUTED)

    # 4. Trust boundary — use the quote layout as the product's central rule.
    slide = prs.slides[3]
    metadata(slide, 4)
    text_box(slide, 0.70, 1.35, 11.6, 3.9,
             "The Evidence Locker\nis the memory.", 48, BLACK, True)
    text_box(slide, 0.75, 5.65, 11.0, 0.42,
             "The LLM is only an interpreter.", 21, TEAL, True)

    # 5. Case workflow — use the text/image split layout.
    slide = prs.slides[4]
    metadata(slide, 5)
    text_box(slide, 0.15, 0.25, 7.0, 1.30,
             "A CASE,\nNOT A CHAT", 31, BLACK, True)
    rich_box(slide, 0.15, 2.05, 7.0, 4.65, [
        ("ASK", 16, BLUE, True, 1),
        ("Hybrid retrieval finds the moments that matter — not the whole archive.", 14, BLACK, False, 13),
        ("CHECK", 16, BLUE, True, 1),
        ("Risk rules decide when the first answer needs a real challenge.", 14, BLACK, False, 13),
        ("RECONCILE", 16, BLUE, True, 1),
        ("Conflicts, uncertainty and Decision Evolution stay visible.", 14, BLACK, False, 13),
        ("RECEIPT", 16, BLUE, True, 1),
        ("Every claim links back to exact stored evidence.", 14, BLACK, False, 0),
    ])

    # 6. Seven-step implementation plan — image-led roadmap layout.
    slide = prs.slides[5]
    metadata(slide, 6)
    roadmap = [
        ("1", "Ingest & parse", "Canonicalize the 45-file archive."),
        ("2", "Lock memory", "2,534 units + stable Evidence IDs."),
        ("3", "Map people", "Attribution, aliases, relationships."),
        ("4", "Find truth", "Hybrid lexical + semantic retrieval."),
        ("5", "Build auditor", "Cases, risk engine, Skeptic."),
        ("6", "Show change", "Decision Evolution + Radar."),
        ("7", "Prove edges", "Pseudonymization, reversal, tests."),
    ]
    for idx, (number, heading, body) in enumerate(roadmap):
        y = 0.52 + idx * 0.58
        text_box(slide, 6.60, y, 0.35, 0.30, number, 13, TEAL, True)
        text_box(slide, 7.02, y, 2.15, 0.25, heading, 13, BLACK, True)
        text_box(slide, 9.20, y, 3.25, 0.28, body, 11, MUTED)
    text_box(slide, 6.58, 5.58, 6.0, 1.65,
             "OUR 7-STEP\nIMPLEMENTATION PLAN", 29, BLACK, True)

    # 7. Differentiators — repurpose the template table as a comparison.
    slide = prs.slides[6]
    metadata(slide, 7)
    text_box(slide, 0.20, 1.15, 4.25, 5.30,
             "Not just\nanother\nchatbot", 36, BLACK, True)
    table = find_table(slide)
    headers = ["Typical chatbot", "Our Auditor", "Why it matters"]
    rows = [
        ["Best-effort prose", "Typed Case + receipt", "Defensible answer"],
        ["Latest statement wins", "Decision Evolution", "Supersession stays visible"],
        ["Waits for a question", "Reconsideration Radar", "Initiative, not just chat"],
    ]
    for col, value in enumerate(headers):
        set_cell(table.cell(0, col), value, 14, True, BLUE)
    for row_idx, row in enumerate(rows, 1):
        for col, value in enumerate(row):
            set_cell(table.cell(row_idx, col), value, 12, False, BLACK)

    # 8. Reversible pseudonymization — another quote layout, now privacy-led.
    slide = prs.slides[7]
    text_box(slide, 6.58, 1.20, 6.35, 4.25,
             '"Privacy changes the\nname in the product —\nnot the history in the archive."',
             34, BLACK, True)
    text_box(slide, 6.60, 5.70, 6.3, 0.32,
             "REVERSIBLE PSEUDONYMIZATION", 15, TEAL, True)
    text_box(slide, 6.60, 6.12, 6.1, 0.28,
             "Preserve · protect · restore", 14, MUTED)

    # 9. Radar initiative — use the image split to make the proactive feature memorable.
    slide = prs.slides[8]
    text_box(slide, 0.18, 0.25, 5.95, 0.65,
             "Rejected/deferred idea → blocker → what changed → assessment",
             14, MUTED, True)
    text_box(slide, 0.18, 1.35, 6.20, 4.55,
             "THE RADAR:\nLOOK AHEAD,\nNOT JUST BACK", 34, BLACK, True)
    text_box(slide, 0.20, 6.20, 6.1, 0.55,
             "7 skeptical checks.\nNo unsupported recommendation.", 16, TEAL, True)

    # 10. Close — retain the template's full-bleed final image.
    slide = prs.slides[9]
    text_box(slide, 0.25, 6.20, 9.75, 0.70,
             "MEMORY WITH A RECEIPT", 32, BLACK, True)
    text_box(slide, 10.35, 6.30, 2.65, 0.64,
             "Auditable memory.\nControlled privacy.", 12, BLACK)

    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
